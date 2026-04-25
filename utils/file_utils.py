import io
import pdfplumber
from docx import Document

def extract_text(file):
    filename = file.filename.lower()
    
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file)
    elif filename.endswith(".txt"):
        return file.read().decode("utf-8")
    else:
        return ""

def extract_text_from_pdf(file):
    text = ""
    # Use pdfplumber for better structure preservation
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text(layout=True)
            if page_text:
                text += page_text + "\n"
    return text

def extract_text_from_docx(file):
    doc = Document(file)
    text = []
    
    # Extract from paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
            
    # Extract from tables
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                text.append(" | ".join(row_data))
                
    return "\n".join(text)
